#!/usr/bin/env python3
"""Self-test for Draftly's helper scripts. Runs in a temporary folder and cleans up.

  selftest.py                 quick test used by setup.sh (3-page PDF -> ingest -> INDEX -> library)
  selftest.py --full          also: PDF without bookmarks (heading split), long PDF (15-page split),
                              .docx, .html, plain text, duplicate check and idempotency
  selftest.py --fixtures DIR  write the generated test files to DIR and keep them

All fixtures are generated here; no copyrighted material is used.
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent


# ---------------------------------------------------------- tiny PDF writer

def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages: list[list[str]]) -> None:
    """Write a minimal text PDF (Helvetica, one line per list item). No dependencies."""
    objs: list[bytes] = []
    n = len(pages)
    font_id = 3
    page_ids = [4 + 2 * i for i in range(n)]
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, lines in enumerate(pages):
        content = ["BT", "/F1 12 Tf", "14 TL", "72 760 Td"]
        for ln in lines:
            content.append(f"({_pdf_escape(ln)}) Tj T*")
        content.append("ET")
        stream = "\n".join(content).encode("latin-1", "replace")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {page_ids[i] + 1} 0 R >>".encode())
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def add_bookmarks_and_labels(path: Path, marks: list[tuple[str, int]], roman_front: int = 0,
                             title: str = "", author: str = "") -> None:
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter(clone_from=PdfReader(str(path)))
    for t, pg in marks:
        writer.add_outline_item(t, pg)
    n = len(writer.pages)
    if roman_front:
        writer.set_page_label(0, roman_front - 1, style="/r")
        writer.set_page_label(roman_front, n - 1, style="/D", start=1)
    if title or author:
        writer.add_metadata({"/Title": title, "/Author": author, "/CreationDate": "D:20240115120000"})
    with open(path, "wb") as f:
        writer.write(f)


def body(topic: str, n: int = 8) -> list[str]:
    return [f"This paragraph discusses {topic} in the context of non-bank financial intermediation,",
            "liquidity mismatch, leverage and the role of central banks as a backstop."] * (n // 2)


def fixtures(d: Path) -> dict[str, Path]:
    d.mkdir(parents=True, exist_ok=True)
    f = {}

    # 1) 3-page PDF with bookmarks
    p = d / "test-bookmarks.pdf"
    make_pdf(p, [["Introduction"] + body("systemic risk"),
                 ["Regulating through banks"] + body("step-in risk"),
                 ["Conclusion"] + body("policy options")])
    add_bookmarks_and_labels(p, [("Introduction", 0), ("Regulating through banks", 1), ("Conclusion", 2)],
                             title="Test Report on Shadow Banking", author="Jane Testauthor")
    f["pdf_bookmarks"] = p

    # 2) 8-page PDF without bookmarks, roman front matter, numbered headings
    p = d / "test-no-bookmarks.pdf"
    make_pdf(p, [["Contents", "1 Introduction ..... 1", "2 Money market funds ..... 3"],
                 ["Foreword"] + body("the foreword"),
                 ["1 Introduction"] + body("the introduction"),
                 body("more introduction"),
                 ["2 Money market funds"] + body("money market fund reform"),
                 body("liquidity fees and gates"),
                 ["3 Conclusions"] + body("the conclusions"),
                 body("final remarks")])
    add_bookmarks_and_labels(p, [], roman_front=2)
    f["pdf_headings"] = p

    # 3) 32-page PDF, no structure -> 15-page chunks
    p = d / "test-long.pdf"
    make_pdf(p, [[f"Running text page {i + 1}"] + body(f"topic {i + 1}") for i in range(32)])
    f["pdf_long"] = p

    # 4) docx with headings and page breaks
    import docx
    from docx.enum.text import WD_BREAK
    doc = docx.Document()
    doc.core_properties.title = "Test Working Paper on LDI Funds"
    doc.core_properties.author = "Alex Testwriter"
    doc.add_heading("Test Working Paper on LDI Funds", 0)
    for i, h in enumerate(["Introduction", "The 2022 gilt market episode", "Conclusion"]):
        doc.add_heading(h, 1)
        doc.add_paragraph(" ".join(body(h.lower(), 4)))
        doc.add_heading(f"{h} - detail", 2)
        para = doc.add_paragraph(" ".join(body("detail", 4)))
        if i < 2:
            para.add_run().add_break(WD_BREAK.PAGE)
    p = d / "test-paper.docx"
    doc.save(str(p))
    f["docx"] = p

    # 5) html
    p = d / "test-article.html"
    p.write_text("""<!doctype html><html><head><title>Test blog on NBFI</title>
<meta name="author" content="Sam Testblogger"></head><body><nav>Home | About</nav>
<article><h1>Test blog on NBFI</h1><h2>Background</h2><p>Non-bank financial intermediation grew after 2008.</p>
<p>It includes money market funds, hedge funds and other investment funds that perform bank-like activities.</p>
<h2>Policy</h2><p>Authorities discuss liquidity management tools and margin practices for these funds.</p>
<ul><li>Liquidity fees</li><li>Swing pricing</li></ul></article><footer>Copyright</footer></body></html>""",
                 encoding="utf-8")
    f["html"] = p

    # 6) markdown
    p = d / "test-notes.md"
    p.write_text("# Test notes\n\nIntro line.\n\n## Part A\n\nText about leverage.\n\n## Part B\n\nText about margins.\n",
                 encoding="utf-8")
    f["md"] = p
    return f


# ------------------------------------------------------------------ helpers

def run(script: str, *args: str) -> dict:
    r = subprocess.run([sys.executable, str(HERE / script), *args], capture_output=True, text=True)
    lines = [ln for ln in r.stdout.strip().splitlines() if ln.strip()]
    if len(lines) != 1:
        raise AssertionError(f"{script} should print exactly one JSON line, got: {r.stdout!r} {r.stderr!r}")
    return json.loads(lines[0])


def check(cond: bool, msg: str, results: list) -> None:
    results.append((cond, msg))
    print(("  PASS  " if cond else "  FAIL  ") + msg)


def make_project(root: Path) -> None:
    for sub in ["references/inbox", "references/pdf", "references/text", "notes", "sessions", "drafts"]:
        (root / sub).mkdir(parents=True, exist_ok=True)
    (root / "references" / "INDEX.md").write_text(
        "# Library index\n\n```\n## template-id\n- Citation: should be ignored\n```\n", encoding="utf-8")


def index_entry(sid: str, citation: str, typ: str, tags: str, relevant: str, status: str, pdf: str) -> str:
    return (f"\n## {sid}\n- Citation: {citation}\n- Type: {typ}\n- Jurisdiction/tags: {tags}\n"
            f"- About: Generated test entry.\n- Key concepts: testing\n- Relevant for: {relevant}\n"
            f"- Status: {status}\n- Files: pdf/{pdf}, text/{sid}/\n")


# -------------------------------------------------------------------- tests

def quick(tmp: Path, results: list) -> None:
    root = tmp / "project"
    make_project(root)
    fx = fixtures(tmp / "fixtures")
    r = run("ingest.py", str(fx["pdf_bookmarks"]), "--id", "testauthor-2024-shadow", "--root", str(root))
    check(r.get("ok") is True, "ingest a 3-page PDF", results)
    check(r.get("split") == "outline" and len(r.get("chapters", [])) == 3, "split by bookmarks into 3 chapters", results)
    text_dir = root / "references" / "text" / "testauthor-2024-shadow"
    first = sorted(text_dir.glob("01-*.md"))
    check(bool(first) and "[p. 1]" in first[0].read_text(), "page markers inserted", results)
    check((text_dir / "_meta.json").exists(), "_meta.json written", results)
    check(fx["pdf_bookmarks"].exists(), "original file left in place (copy mode)", results)
    check((root / "references" / "pdf" / "testauthor-2024-shadow.pdf").exists(), "original copied into the library", results)

    idx = root / "references" / "INDEX.md"
    idx.write_text(idx.read_text() + index_entry(
        "testauthor-2024-shadow", "Jane Testauthor, *Test Report on Shadow Banking* (Test Press 2024)",
        "report", "UK, theory, maybe", "2.1, 3 (pp 1-3)", "skimmed", "testauthor-2024-shadow.pdf"), encoding="utf-8")
    (root / "drafts" / "ch1-test-v1.md").write_text(
        "Claim one.[CITATION NEEDED: FSB report] Claim two [PINPOINT NEEDED].\n", encoding="utf-8")
    (root / "thesis-brief.md").write_text(
        "# Thesis brief\n- Working title: Test thesis\n\n## Outline\n- 1 Introduction\n- 2.1 Banks\n- 3 Funds\n",
        encoding="utf-8")
    b = run("build_library.py", "--root", str(root))
    html = (root / "library.html").read_text() if (root / "library.html").exists() else ""
    check(b.get("ok") is True and b.get("sources") == 1, "build library.html from INDEX.md", results)
    check("testauthor-2024-shadow" in html and "template-id" not in html, "entry rendered, template ignored", results)
    check(b.get("citation_needed") == 1 and b.get("pinpoint_needed") == 1, "citation markers counted", results)
    check(b.get("sections_without_sources") == ["1"], "sections without sources detected", results)
    check("http://" not in html.replace("http://www.w3.org", "") and "https://" not in html and "<link" not in html,
          "library.html is self-contained (no external resources)", results)


def full(tmp: Path, results: list) -> None:
    root = tmp / "project-full"
    make_project(root)
    fx = fixtures(tmp / "fixtures-full")
    r = run("ingest.py", str(fx["pdf_headings"]), "--id", "test-2024-mmf", "--root", str(root))
    titles = [c["title"] for c in r.get("chapters", [])]
    check(r.get("ok") and r.get("split") == "headings", "PDF without bookmarks split by detected headings", results)
    check("1 Introduction" in titles and "2 Money market funds" in titles, f"headings found: {titles}", results)
    txt = "".join(p.read_text() for p in sorted((root / "references/text/test-2024-mmf").glob("*.md")))
    check("[p. i]" in txt and "[p. 1]" in txt, "printed page labels used (roman + arabic)", results)

    r = run("ingest.py", str(fx["pdf_long"]), "--id", "test-2024-long", "--root", str(root))
    check(r.get("ok") and r.get("split") == "fixed-15-pages" and len(r["chapters"]) == 3, "long PDF split every 15 pages", results)

    r = run("ingest.py", str(fx["docx"]), "--id", "testwriter-2024-ldi", "--root", str(root))
    check(r.get("ok") and len(r.get("chapters", [])) == 3, f"docx split by Heading 1 ({len(r.get('chapters', []))} parts)", results)
    txt = "".join(p.read_text() for p in sorted((root / "references/text/testwriter-2024-ldi").glob("*.md")))
    check("[p. 2]" in txt and "[p. 3]" in txt, "docx page breaks become page markers", results)
    check(r.get("author") == "Alex Testwriter", "docx author detected", results)

    r = run("ingest.py", str(fx["html"]), "--id", "testblogger-2024-nbfi", "--root", str(root))
    check(r.get("ok") and r.get("title") == "Test blog on NBFI" and len(r.get("chapters", [])) == 2, "HTML extracted and split", results)
    txt = "".join(p.read_text() for p in sorted((root / "references/text/testblogger-2024-nbfi").glob("*.md")))
    check("Home | About" not in txt and "Swing pricing" in txt, "HTML boilerplate removed, content kept", results)

    r = run("ingest.py", str(fx["md"]), "--id", "test-notes", "--root", str(root))
    check(r.get("ok") and len(r.get("chapters", [])) >= 2, "markdown split by headings", results)

    r = run("ingest.py", str(fx["pdf_long"]), "--id", "test-2024-long", "--root", str(root))
    check(r.get("ok") and r.get("status") == "unchanged", "re-ingest is idempotent (unchanged)", results)
    r = run("ingest.py", str(fx["pdf_long"]), "--id", "another-id", "--root", str(root))
    check(not r.get("ok") and r.get("duplicate_of") == "test-2024-long", "duplicate under another id is refused", results)
    r = run("ingest.py", str(fx["pdf_long"]), "--check-only", "--root", str(root))
    check(r.get("duplicate_of") == "test-2024-long", "--check-only reports the duplicate", results)

    scan = tmp / "scan.pdf"
    make_pdf(scan, [[""], [""]])
    r = run("ingest.py", str(scan), "--id", "test-scan", "--root", str(root))
    check(r.get("ok") or r.get("error") == "scanned_no_ocr", "scanned PDF: OCR or a clear message", results)

    r = run("ingest.py", str(tmp / "missing.pdf"), "--root", str(root))
    check(not r.get("ok") and r.get("error") == "not_found", "missing file gives a clear error", results)

    inbox_copy = root / "references" / "inbox" / "dropped.pdf"
    shutil.copy2(fx["pdf_bookmarks"], inbox_copy)
    r = run("ingest.py", str(inbox_copy), "--id", "test-inbox", "--store", "move", "--root", str(root))
    check(r.get("ok") and not inbox_copy.exists() and (root / "references/pdf/test-inbox.pdf").exists(),
          "inbox file moved into the library with --store move", results)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--full", action="store_true")
    ap.add_argument("--fixtures", help="write fixtures to this folder and exit")
    args = ap.parse_args()
    if args.fixtures:
        for k, p in fixtures(Path(args.fixtures)).items():
            print(f"{k}: {p}")
        return 0
    results: list = []
    tmp = Path(tempfile.mkdtemp(prefix="draftly-selftest-"))
    try:
        print("Draftly self-test")
        quick(tmp, results)
        if args.full:
            full(tmp, results)
    except Exception as e:  # noqa: BLE001
        results.append((False, f"crashed: {type(e).__name__}: {e}"))
        print(f"  FAIL  crashed: {type(e).__name__}: {e}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    failed = [m for ok, m in results if not ok]
    print(f"{len(results) - len(failed)}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
